import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def add_thank_you_slide(ppt_path="GPON_Network_Mapping_Project_Report.pptx"):
    prs = Presentation(ppt_path)
    blank_layout = prs.slide_layouts[6]
    
    # Check if a Thank You slide is already present at the end
    last_slide = prs.slides[-1]
    has_thank_you = False
    for shape in last_slide.shapes:
        if shape.has_text_frame and "thank you" in shape.text_frame.text.lower():
            has_thank_you = True
            break
            
    if has_thank_you:
        print("Thank You slide already exists at the end. Skipping duplicate.")
        return

    # Colors
    DARK_NAVY = RGBColor(15, 23, 42)      # #0F172A
    CYAN_BLUE = RGBColor(2, 132, 199)     # #0284C7
    EMERALD = RGBColor(16, 185, 129)      # #10B981
    WHITE = RGBColor(255, 255, 255)
    SLATE_LIGHT = RGBColor(226, 232, 240) # #E2E8F0
    SLATE_MUTED = RGBColor(148, 163, 184) # #94A3B8

    slide = prs.slides.add_slide(blank_layout)

    # 1. Dark Navy Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_NAVY
    bg.line.fill.background()

    # 2. Top Cyan Accent Strip
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = CYAN_BLUE
    top_bar.line.fill.background()

    # 3. App Emblem Icon (Right side)
    icon_path = os.path.join("web_app", "icon-1024.png")
    if os.path.exists(icon_path):
        slide.shapes.add_picture(icon_path, Inches(9.0), Inches(1.8), Inches(3.2), Inches(3.2))

    # 4. Text Content Box (Left side)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(7.6), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Badge
    p_badge = tf.paragraphs[0]
    p_badge.text = "FIBER & GPON NETWORK MAPPING SYSTEM"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = EMERALD

    # Thank You Title
    p_title = tf.add_paragraph()
    p_title.text = "Thank You!"
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_before = Pt(6)

    # Subtitle
    p_sub = tf.add_paragraph()
    p_sub.text = "Open for Questions, Technical Discussion & Demonstration"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = SLATE_MUTED
    p_sub.space_before = Pt(8)

    # Accent divider line
    p_div = tf.add_paragraph()
    p_div.text = "────────────────────────────────────────"
    p_div.font.size = Pt(8)
    p_div.font.color.rgb = RGBColor(51, 65, 85)
    p_div.space_before = Pt(6)

    # Highlights
    bullets = [
        "Interactive GIS Mobile Survey PWA (Zero App-Store friction, Offline-First)",
        "Sub-meter GPS spatial accuracy for KSEB electric poles & optical nodes",
        "Role-Based Access Governance across 5 administrative circle tiers",
        "Central Office Management Portal with instant Excel & ZIP export streaming",
        "Automated 24/7 Google Drive disaster recovery with real-time health telemetry"
    ]
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = "✓  " + b
        bp.font.size = Pt(12)
        bp.font.color.rgb = SLATE_LIGHT
        bp.space_before = Pt(7)

    # 5. Bottom Metadata Card
    foot_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(6.45), Inches(11.333), Inches(0.65))
    foot_box.fill.solid()
    foot_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
    foot_box.line.color.rgb = RGBColor(51, 65, 85)

    ftf = foot_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_right = Inches(0.25)
    ftf.margin_top = Inches(0.12)

    fp = ftf.paragraphs[0]
    fp.text = "Submitted to: Superior Officers & Circle Leadership | Kerala Telecom Operations | Production Ready"
    fp.font.size = Pt(11)
    fp.font.bold = True
    fp.font.color.rgb = WHITE

    prs.save(ppt_path)
    print(f"Thank You slide added successfully to: {ppt_path}")

    # Copy to web_app and artifact directory so the admin portal serves the latest version
    web_app_ppt = os.path.join("web_app", "GPON_Network_Mapping_Project_Report.pptx")
    try:
        shutil.copy2(ppt_path, web_app_ppt)
        print(f"Copied to web_app: {web_app_ppt}")
    except Exception as e:
        print(f"Warning: {e}")

    artifact_dir = r"C:\Users\Anoop P\.gemini\antigravity\brain\80fdf803-a737-43de-9989-3958386366af"
    if os.path.exists(artifact_dir):
        try:
            shutil.copy2(ppt_path, os.path.join(artifact_dir, "GPON_Network_Mapping_Project_Report.pptx"))
            print(f"Copied to artifact dir: {artifact_dir}")
        except Exception as e:
            print(f"Warning: {e}")

if __name__ == "__main__":
    add_thank_you_slide()
